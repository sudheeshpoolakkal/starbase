import sys
import os
import subprocess
from pathlib import Path
import pypdf

# Add current directory to path
sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from generate_academic_report import (
    ROOT_DIR, DOCX_OUT, PDF_OUT
)
from report_front_matter import (
    add_cover_page, add_certificate, add_declaration, add_abstract,
    add_toc, add_list_of_figures, add_list_of_tables
)
from report_ch1_ch2 import add_chapter_1, add_chapter_2
from report_ch3 import add_chapter_3
from report_ch4_ch5_ch6 import (
    add_chapter_4, add_chapter_5, add_chapter_6, add_bibliography
)

def build_full_report(toc_map=None, lof_map=None, lot_map=None):
    doc = Document()

    # SECTION 1: Cover Page (No footer / no page number)
    sec1 = doc.sections[0]
    sec1.page_width = Inches(8.2677)
    sec1.page_height = Inches(11.6929)
    sec1.left_margin = Inches(1.25)
    sec1.right_margin = Inches(1.0)
    sec1.top_margin = Inches(1.0)
    sec1.bottom_margin = Inches(1.0)
    add_cover_page(doc)

    # SECTION 2: Prelims (Certificate to List of Tables) with lowerRoman numbering
    sec2 = doc.add_section()
    sec2.page_width = Inches(8.2677)
    sec2.page_height = Inches(11.6929)
    sec2.left_margin = Inches(1.25)
    sec2.right_margin = Inches(1.0)
    sec2.top_margin = Inches(1.0)
    sec2.bottom_margin = Inches(1.0)
    sec2.footer.is_linked_to_previous = False

    pgNumType2 = OxmlElement('w:pgNumType')
    pgNumType2.set(qn('w:fmt'), 'lowerRoman')
    pgNumType2.set(qn('w:start'), '1')
    sec2._sectPr.append(pgNumType2)

    p_foot2 = sec2.footer.paragraphs[0]
    p_foot2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_fld2 = OxmlElement('w:fldSimple')
    r_fld2.set(qn('w:instr'), 'PAGE')
    p_foot2._p.append(r_fld2)

    add_certificate(doc)
    add_declaration(doc)
    add_abstract(doc)

    # TOC Entries
    def t(label, default="1"):
        return toc_map.get(label, default) if toc_map else default
    def f(label, default="1"):
        return lof_map.get(label, default) if lof_map else default
    def tb(label, default="1"):
        return lot_map.get(label, default) if lot_map else default

    toc_entries = [
        ("CERTIFICATE OF COMPLETION", t("CERTIFICATE OF COMPLETION", "i"), 0),
        ("DECLARATION", t("DECLARATION", "ii"), 0),
        ("ABSTRACT", t("ABSTRACT", "iii"), 0),
        ("LIST OF FIGURES", t("LIST OF FIGURES", "vi"), 0),
        ("LIST OF TABLES", t("LIST OF TABLES", "vii"), 0),
        ("1 Introduction", t("1 INTRODUCTION", "1"), 0),
        ("1.1 Introduction", t("1.1 Introduction", "1"), 1),
        ("1.2 Problem Statement", t("1.2 Problem Statement", "1"), 1),
        ("1.3 Objectives", t("1.3 Objectives", "2"), 1),
        ("1.4 Scope of the Project", t("1.4 Scope of the Project", "2"), 1),
        ("2 System Analysis and Design", t("2 SYSTEM ANALYSIS AND DESIGN", "4"), 0),
        ("2.1 Existing System and Its Limitations", t("2.1 Existing System and Its Limitations", "4"), 1),
        ("2.2 Proposed System", t("2.2 Proposed System", "4"), 1),
        ("2.3 System Requirements", t("2.3 System Requirements", "5"), 1),
        ("2.4 System Architecture", t("2.4 System Architecture", "6"), 1),
        ("2.5 System Workflow", t("2.5 System Workflow", "7"), 1),
        ("2.6 Dataflow Diagrams", t("2.6 Dataflow Diagrams", "8"), 1),
        ("3 Database Design", t("3 DATABASE DESIGN", "10"), 0),
        ("3.1 Database Overview", t("3.1 Database Overview", "10"), 1),
        ("3.2 Three-Schema Architecture", t("3.2 Three-Schema Architecture", "10"), 1),
        ("3.3 Entity–Relationship (ER) Diagram", t("3.3 Entity–Relationship (ER) Diagram", "11"), 1),
        ("3.4 Relational Schema & Normalization Analysis", t("3.4 Relational Schema & Normalization Analysis", "12"), 1),
        ("3.5 Database Tables and Attributes", t("3.5 Database Tables and Attributes", "14"), 1),
        ("3.6 Primary Keys, Foreign Keys and Constraints", t("3.6 Primary Keys, Foreign Keys and Constraints", "16"), 1),
        ("4 Implementation", t("4 IMPLEMENTATION", "18"), 0),
        ("4.1 Technologies Used", t("4.1 Technologies Used", "18"), 1),
        ("4.2 Module Description", t("4.2 Module Description", "18"), 1),
        ("4.3 Database Implementation", t("4.3 Database Implementation", "19"), 1),
        ("4.4 SQL Queries, Operations, and Scalability Evidence", t("4.4 SQL Queries, Operations, and Scalability Evidence", "20"), 1),
        ("4.5 Application Implementation", t("4.5 Application Implementation", "22"), 1),
        ("4.6 Database Connectivity", t("4.6 Database Connectivity", "22"), 1),
        ("4.7 Challenges Faced During the Project", t("4.7 Challenges Faced During the Project", "23"), 1),
        ("5 Testing and Results", t("5 TESTING AND RESULTS", "24"), 0),
        ("5.1 Testing Methodology", t("5.1 Testing Methodology", "24"), 1),
        ("5.2 Test Cases and Results", t("5.2 Test Cases and Results", "24"), 1),
        ("5.3 Application Screenshots", t("5.3 Application Screenshots", "26"), 1),
        ("5.4 Results and Discussion", t("5.4 Results and Discussion", "29"), 1),
        ("6 Conclusion", t("6 CONCLUSION", "30"), 0),
        ("6.1 Conclusion", t("6.1 Conclusion", "30"), 1),
        ("6.2 Limitations", t("6.2 Limitations", "30"), 1),
        ("6.3 Future Enhancements", t("6.3 Future Enhancements", "31"), 1),
        ("BIBLIOGRAPHY", t("BIBLIOGRAPHY", "32"), 0),
    ]
    add_toc(doc, toc_entries)

    lof_entries = [
        ("Figure 2.1 — Starbase three-tier system architecture and offline ingestion pipeline", f("Figure 2.1", "6")),
        ("Figure 2.2 — Level 0 Data Flow Diagram (Context Analysis)", f("Figure 2.2", "8")),
        ("Figure 2.3 — Level 1 Data Flow Diagram (Decomposed Functional Modules)", f("Figure 2.3", "9")),
        ("Figure 3.1 — Entity–Relationship (ER) diagram for Starbase material catalog", f("Figure 3.1", "11")),
        ("Figure 5.1 — Starbase Desktop Landing and Overview Interface", f("Figure 5.1", "26")),
        ("Figure 5.2 — Responsive Mobile Material Exploration Interface", f("Figure 5.2", "27")),
        ("Figure 5.3 — Material Library with Property-Value Sorting and Provenance Tracking", f("Figure 5.3", "27")),
        ("Figure 5.4 — Multi-Attribute Selection Studio with Ranked Recommendations", f("Figure 5.4", "28")),
    ]
    add_list_of_figures(doc, lof_entries)

    lot_entries = [
        ("Table 2.1 — System Requirements Specifications", tb("Table 2.1", "5")),
        ("Table 2.2 — DFD Process Specifications and Data Store Interactions", tb("Table 2.2", "8")),
        ("Table 3.1 — Database Schema Architecture and Entity Groupings", tb("Table 3.1", "10")),
        ("Table 3.2 — Three-Schema Architecture Implementation in Starbase", tb("Table 3.2", "11")),
        ("Table 3.3 — Entity Cardinality and Relationship Constraints", tb("Table 3.3", "12")),
        ("Table 3.4 — Data Dictionary for Data_Source Entity", tb("Table 3.4", "14")),
        ("Table 3.5 — Data Dictionary for Material Entity", tb("Table 3.5", "14")),
        ("Table 3.6 — Data Dictionary for Property Entity", tb("Table 3.6", "15")),
        ("Table 3.7 — Data Dictionary for Material_Property Association Table", tb("Table 3.7", "15")),
        ("Table 3.16 — Relational Constraints and Cascade Actions", tb("Table 3.16", "16")),
        ("Table 3.17 — Database Indexing Strategy and Query Optimization Roles", tb("Table 3.17", "17")),
        ("Table 4.1 — Software Technology Stack and Architectural Roles", tb("Table 4.1", "18")),
        ("Table 4.2 — Software Module Responsibilities and Source Code Mapping", tb("Table 4.2", "18")),
        ("Table 4.3 — Ingested Catalog Data Sources and Ingestion Metrics", tb("Table 4.3", "19")),
        ("Table 4.4 — Starbase RESTful API Endpoint Specifications", tb("Table 4.4", "21")),
        ("Table 4.7 — Measured Query Execution Benchmarks and EXPLAIN Plans", tb("Table 4.7", "21")),
        ("Table 5.1 — Automated Integration and Regression Test Suite Results", tb("Table 5.1", "25")),
        ("Table 5.2 — Deployment and Hosting Architecture Status Matrix", tb("Table 5.2", "29")),
    ]
    add_list_of_tables(doc, lot_entries)

    # SECTION 3: Body (Chapters 1 to 6 + Bibliography) with decimal numbering starting at 1
    sec3 = doc.add_section()
    sec3.page_width = Inches(8.2677)
    sec3.page_height = Inches(11.6929)
    sec3.left_margin = Inches(1.25)
    sec3.right_margin = Inches(1.0)
    sec3.top_margin = Inches(1.0)
    sec3.bottom_margin = Inches(1.0)
    sec3.footer.is_linked_to_previous = False

    pgNumType3 = OxmlElement('w:pgNumType')
    pgNumType3.set(qn('w:fmt'), 'decimal')
    pgNumType3.set(qn('w:start'), '1')
    sec3._sectPr.append(pgNumType3)

    p_foot3 = sec3.footer.paragraphs[0]
    p_foot3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_fld3 = OxmlElement('w:fldSimple')
    r_fld3.set(qn('w:instr'), 'PAGE')
    p_foot3._p.append(r_fld3)

    # Chapters
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_chapter_5(doc)
    add_chapter_6(doc)
    add_bibliography(doc)

    return doc

def extract_pdf_page_maps(pdf_path):
    reader = pypdf.PdfReader(str(pdf_path))
    num_pages = len(reader.pages)
    print(f"Analyzing PDF: {num_pages} total pages")

    # Locate body start index
    body_start_idx = None
    for idx, page in enumerate(reader.pages):
        lines = page.extract_text().splitlines()
        if any(l.strip() == "1 INTRODUCTION" for l in lines):
            body_start_idx = idx
            break

    if body_start_idx is None:
        body_start_idx = 8

    toc_map = {
        "CERTIFICATE OF COMPLETION": "i",
        "DECLARATION": "ii",
        "ABSTRACT": "iii",
        "LIST OF FIGURES": "vi",
        "LIST OF TABLES": "vii"
    }
    lof_map = {}
    lot_map = {}

    target_headings = [
        "1 INTRODUCTION", "1.1 Introduction", "1.2 Problem Statement",
        "1.3 Objectives", "1.4 Scope of the Project",
        "2 SYSTEM ANALYSIS AND DESIGN", "2.1 Existing System and Its Limitations",
        "2.2 Proposed System", "2.3 System Requirements",
        "2.4 System Architecture", "2.5 System Workflow",
        "2.6 Dataflow Diagrams",
        "3 DATABASE DESIGN", "3.1 Database Overview",
        "3.2 Three-Schema Architecture", "3.3 Entity–Relationship (ER) Diagram",
        "3.4 Relational Schema & Normalization Analysis",
        "3.5 Database Tables and Attributes",
        "3.6 Primary Keys, Foreign Keys and Constraints",
        "4 IMPLEMENTATION", "4.1 Technologies Used",
        "4.2 Module Description", "4.3 Database Implementation",
        "4.4 SQL Queries, Operations, and Scalability Evidence",
        "4.5 Application Implementation", "4.6 Database Connectivity",
        "4.7 Challenges Faced During the Project",
        "5 TESTING AND RESULTS", "5.1 Testing Methodology",
        "5.2 Test Cases and Results", "5.3 Application Screenshots",
        "5.4 Results and Discussion",
        "6 CONCLUSION", "6.1 Conclusion", "6.2 Limitations",
        "6.3 Future Enhancements", "BIBLIOGRAPHY"
    ]

    for idx in range(body_start_idx, num_pages):
        arabic_page = str(idx - body_start_idx + 1)
        text = reader.pages[idx].extract_text()
        for line in text.splitlines():
            l_s = line.strip()
            for h in target_headings:
                if h not in toc_map and (l_s == h or (len(h) > 8 and h in l_s)):
                    toc_map[h] = arabic_page
            if l_s.startswith("Figure ") and ("—" in l_s or "-" in l_s):
                fig_key = l_s.split("—")[0].split("-")[0].strip()
                if fig_key not in lof_map:
                    lof_map[fig_key] = arabic_page
            elif l_s.startswith("Table ") and ("—" in l_s or "-" in l_s):
                tbl_key = l_s.split("—")[0].split("-")[0].strip()
                if tbl_key not in lot_map:
                    lot_map[tbl_key] = arabic_page

    print(f"Extracted {len(toc_map)} TOC entries, {len(lof_map)} Figures, {len(lot_map)} Tables.")
    return toc_map, lof_map, lot_map

def main():
    print("=== PASS 1: Generating Initial Document ===")
    pass1_docx = Path("/tmp/starbase_pass1.docx")
    pass1_pdf = Path("/tmp/starbase_pass1.pdf")
    doc1 = build_full_report()
    doc1.save(str(pass1_docx))
    print(f"Saved Pass 1 DOCX: {pass1_docx}")

    print("Compiling Pass 1 PDF with LibreOffice...")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", str(pass1_docx), "--outdir", "/tmp"], check=True)
    
    print("Extracting exact page numbers from PDF...")
    toc_map, lof_map, lot_map = extract_pdf_page_maps(pass1_pdf)

    print("=== PASS 2: Rebuilding Document with True Synchronized Page Numbers ===")
    doc2 = build_full_report(toc_map, lof_map, lot_map)
    doc2.save(str(DOCX_OUT))
    print(f"Saved Final Academic DOCX: {DOCX_OUT}")

    print("Compiling Final PDF...")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", str(DOCX_OUT), "--outdir", str(ROOT_DIR)], check=True)
    print(f"Saved Final Academic PDF: {PDF_OUT}")

    final_reader = pypdf.PdfReader(str(PDF_OUT))
    print(f"SUCCESS: Final document generated with {len(final_reader.pages)} pages.")

if __name__ == "__main__":
    main()
