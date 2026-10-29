from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from generate_academic_report import (
    CREST_PATH, add_p, add_leader_line
)

def add_cover_page(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.2677)
    sec.page_height = Inches(11.6929)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)

    # Crest / Emblem at top
    if CREST_PATH.exists():
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(0)
        p_logo.paragraph_format.space_after = Pt(20)
        p_logo.add_run().add_picture(str(CREST_PATH), width=Inches(1.6))

    # PROJECT TITLE (Bold, Italic, 22pt, Uppercase)
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(16)
    p_title.paragraph_format.line_spacing = Pt(26)
    r_title = p_title.add_run("STARBASE MATERIAL SELECTION\nSYSTEM")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(22)
    r_title.bold = True
    r_title.italic = True

    # PROJECT REPORT (Bold, 16pt, Uppercase)
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("PROJECT REPORT")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(16)
    r_sub.bold = True

    # Submitted by (Italic, 12pt)
    p_by = doc.add_paragraph()
    p_by.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_by.paragraph_format.space_before = Pt(0)
    p_by.paragraph_format.space_after = Pt(16)
    r_by = p_by.add_run("Submitted by")
    r_by.font.name = 'Times New Roman'
    r_by.font.size = Pt(12)
    r_by.italic = True

    # Members list with Sudheesh P on top
    members = [
        ("SUDHEESH P", "64", "LTVE24EL071"),
        ("BABITHOSH EB", "18", "LTVE24EL069"),
        ("MINNA ANTONY", "47", "TVE24EL044"),
        ("SEBIN XAVIER", "60", "TVE24EL056")
    ]

    for name, roll, uid in members:
        p_m = doc.add_paragraph()
        p_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_m.paragraph_format.space_before = Pt(0)
        p_m.paragraph_format.space_after = Pt(22)
        p_m.paragraph_format.line_spacing = Pt(20)

        r_n = p_m.add_run(f"{name}\n")
        r_n.font.name = 'Times New Roman'
        r_n.font.size = Pt(12)
        r_n.bold = True

        r_r = p_m.add_run(f"Roll No: {roll}\n")
        r_r.font.name = 'Times New Roman'
        r_r.font.size = Pt(12)
        r_r.bold = True

        r_u = p_m.add_run(f"University ID: {uid}")
        r_u.font.name = 'Times New Roman'
        r_u.font.size = Pt(12)
        r_u.bold = True

    # Affiliation block
    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_before = Pt(28)
    p_dept.paragraph_format.space_after = Pt(28)
    p_dept.paragraph_format.line_spacing = Pt(20)

    r_deg = p_dept.add_run("B.Tech Electrical and Computer Engineering\n")
    r_deg.font.name = 'Times New Roman'
    r_deg.font.size = Pt(12)
    r_deg.bold = True

    r_dp = p_dept.add_run("Department of Electrical Engineering\n")
    r_dp.font.name = 'Times New Roman'
    r_dp.font.size = Pt(12)
    r_dp.bold = True

    r_col = p_dept.add_run("College of Engineering Trivandrum")
    r_col.font.name = 'Times New Roman'
    r_col.font.size = Pt(12)
    r_col.bold = True

    # Academic Year
    p_yr = doc.add_paragraph()
    p_yr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_yr.paragraph_format.space_before = Pt(0)
    p_yr.paragraph_format.space_after = Pt(0)
    r_yr = p_yr.add_run("Academic Year 2026–2027")
    r_yr.font.name = 'Times New Roman'
    r_yr.font.size = Pt(12)
    r_yr.bold = True

def add_certificate(doc):
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(12)
    p_h.paragraph_format.space_after = Pt(18)
    r_h = p_h.add_run("CERTIFICATE OF COMPLETION")
    r_h.font.name = 'Times New Roman'
    r_h.font.size = Pt(16)
    r_h.bold = True

    add_p(doc, "This is to certify that Sudheesh P (Roll No. 64), Babithosh EB (Roll No. 18), Minna Antony (Roll No. 47), and Sebin Xavier (Roll No. 60), B.Tech Electrical and Computer Engineering students of College of Engineering Trivandrum, have completed the work presented in the report titled “Starbase Material Selection System” at College of Engineering Trivandrum, as part of the Database Management System (PBEOT504) Project-Based Learning (PBL) curriculum. This is a bonafide record of their original work carried out under our supervision and guidance.", space_after=24)

    # Guides Signatures
    p_sig1 = doc.add_paragraph()
    p_sig1.paragraph_format.space_before = Pt(36)
    p_sig1.paragraph_format.space_after = Pt(2)
    p_sig1.paragraph_format.line_spacing = 1.15
    r1 = p_sig1.add_run("Prof. Muhammed Raseem S.\n")
    r1.bold = True
    r1.font.name = 'Times New Roman'; r1.font.size = Pt(11.5)
    r2 = p_sig1.add_run("Department of Computer Science & Engineering\nCollege of Engineering Trivandrum\nSignature: __________________________    Date: ______________")
    r2.font.name = 'Times New Roman'; r2.font.size = Pt(11)

    p_sig2 = doc.add_paragraph()
    p_sig2.paragraph_format.space_before = Pt(28)
    p_sig2.paragraph_format.space_after = Pt(2)
    p_sig2.paragraph_format.line_spacing = 1.15
    r3 = p_sig2.add_run("Prof. Julie T.S.\n")
    r3.bold = True
    r3.font.name = 'Times New Roman'; r3.font.size = Pt(11.5)
    r4 = p_sig2.add_run("Department of Electrical Engineering\nCollege of Engineering Trivandrum\nSignature: __________________________    Date: ______________")
    r4.font.name = 'Times New Roman'; r4.font.size = Pt(11)

def add_declaration(doc):
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(12)
    p_h.paragraph_format.space_after = Pt(18)
    p_h.paragraph_format.page_break_before = True
    r_h = p_h.add_run("DECLARATION")
    r_h.font.name = 'Times New Roman'
    r_h.font.size = Pt(16)
    r_h.bold = True

    add_p(doc, "We hereby declare that the project report entitled “Starbase Material Selection System” is a record of the work carried out by our group as part of the Project-Based Learning (PBL) course Database Management System (PBEOT504) under the guidance of Prof. Muhammed Raseem S. and Prof. Julie T.S.", space_after=8)
    add_p(doc, "We further declare that this report has been prepared as part of the academic requirements of the B.Tech degree programme and that the work presented herein has not been submitted, either in whole or in part, to any other institution for the award of any degree or diploma.", space_after=8)
    add_p(doc, "We have duly acknowledged all sources of information used in the preparation of this report.", space_after=20)

    # Student Signatures with Sudheesh on top
    p_place = doc.add_paragraph()
    p_place.paragraph_format.space_before = Pt(12)
    p_place.paragraph_format.space_after = Pt(16)
    r_pl = p_place.add_run("TRIVANDRUM\n26-10-2026")
    r_pl.font.name = 'Times New Roman'; r_pl.font.size = Pt(11.5); r_pl.bold = True

    p_sigs = doc.add_paragraph()
    p_sigs.paragraph_format.space_before = Pt(12)
    p_sigs.paragraph_format.line_spacing = 1.6
    sigs = [
        "Sudheesh P   (Roll No. 64, TVE ID: LTVE24EL071)     Signature: ___________________",
        "Babithosh EB (Roll No. 18, TVE ID: LTVE24EL069)     Signature: ___________________",
        "Minna Antony (Roll No. 47, TVE ID: TVE24EL044)     Signature: ___________________",
        "Sebin Xavier (Roll No. 60, TVE ID: TVE24EL056)     Signature: ___________________"
    ]
    for s in sigs:
        r = p_sigs.add_run(s + "\n")
        r.font.name = 'Times New Roman'; r.font.size = Pt(10.5)

def add_abstract(doc):
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(12)
    p_h.paragraph_format.space_after = Pt(18)
    p_h.paragraph_format.page_break_before = True
    r_h = p_h.add_run("ABSTRACT")
    r_h.font.name = 'Times New Roman'
    r_h.font.size = Pt(16)
    r_h.bold = True

    add_p(doc, "Selecting an optimal engineering material for a mechanical or aerospace component requires balancing conflicting design objectives, including mass minimization, structural yield strength, stiffness, thermal endurance, and commercial availability. While vast computational and empirical materials data exists, it remains fragmented across incompatible repositories, proprietary silos, and non-standardized datasheets. Starbase is an open, relational database system and engineering selection workspace designed to unify multi-source scientific materials data into a single, fully normalized catalog.", space_after=8)

    add_p(doc, "The database schema implements a third-normal-form (3NF) property-dictionary architecture that reconciles 6,265,691 material records across 14 authoritative open repositories, including NIST JARVIS-DFT, Materials Project (core and GNoME), AFLOW, NOMAD, Alexandria, NASA TPSX, spacecraft outgassing archives, high-entropy alloy datasets, and certified stainless-steel supplier catalogs. Complete source provenance (origin repository, upstream version, license terms, digital checksum, and evidence classification) is captured for every record. The database footprint comprises 18 core application schema tables (17 active application tables and one legacy review table), expanding to 20 tables in the cloud deployment with two dedicated migration and audit support tables.", space_after=8)

    add_p(doc, "The selection engine applies a hybrid evaluation strategy: mandatory engineering constraints are evaluated directly in SQL via composite-indexed criteria queries, followed by multi-attribute min-max normalized weighted scoring that ranks surviving candidates. Crucially, missing values are never zero-imputed; candidates lacking required properties are explicitly disqualified with explainable diagnostic logs. Starbase includes a React 18 single-page application, an Express 5 REST API, and dual database support across local MySQL 8.4 and TiDB Cloud Serverless. An automated test suite of 29 test suites verifies ranking algebra, transactional persistence, role-based access control, and large-scale catalog pagination.", space_after=0)

def add_toc(doc, toc_entries):
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(12)
    p_h.paragraph_format.space_after = Pt(16)
    p_h.paragraph_format.page_break_before = True
    r_h = p_h.add_run("CONTENTS")
    r_h.font.name = 'Times New Roman'
    r_h.font.size = Pt(16)
    r_h.bold = True

    for label, page, level in toc_entries:
        add_leader_line(doc, label, page, level=level, is_bold=(level == 0))

def add_list_of_figures(doc, figure_entries):
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(12)
    p_h.paragraph_format.space_after = Pt(16)
    p_h.paragraph_format.page_break_before = True
    r_h = p_h.add_run("LIST OF FIGURES")
    r_h.font.name = 'Times New Roman'
    r_h.font.size = Pt(16)
    r_h.bold = True

    for label, page in figure_entries:
        add_leader_line(doc, label, page, level=0, is_bold=False)

def add_list_of_tables(doc, table_entries):
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(12)
    p_h.paragraph_format.space_after = Pt(16)
    p_h.paragraph_format.page_break_before = True
    r_h = p_h.add_run("LIST OF TABLES")
    r_h.font.name = 'Times New Roman'
    r_h.font.size = Pt(16)
    r_h.bold = True

    for label, page in table_entries:
        add_leader_line(doc, label, page, level=0, is_bold=False)
