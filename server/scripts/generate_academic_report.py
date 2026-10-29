#!/usr/bin/env python3
"""
Starbase Material Selection System — Academic Project Report Generator
Conforms to CET B.Tech PBL Database Management System (PBEOT504) Report Outline.
Generates an academic report with embedded diagrams, screenshots,
in-text citations, 3NF normalization proofs, EXPLAIN scalability benchmarks,
and verified table counts.
"""

import os
import sys
import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT_DIR = Path(__file__).resolve().parent.parent.parent
DOCX_OUT = ROOT_DIR / "Starbase Material Selection System — Project Report.docx"
PDF_OUT = ROOT_DIR / "Starbase Material Selection System — Project Report.pdf"

CREST_PATH = ROOT_DIR / "docs/report-screenshots/cet-crest.png"
DIAGRAMS_DIR = ROOT_DIR / "docs/diagram-renders"
SCREENSHOTS_DIR = ROOT_DIR / "docs/report-screenshots"

# Helper XML styling
def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)

def set_table_borders(table, color="B0B8C1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def style_table(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trPr.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                set_cell_background(cell, "EAEFF5")
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.line_spacing = 1.05
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.bold = True
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(9.5)
        else:
            bg = "FAFBFC" if i % 2 == 1 else "FFFFFF"
            for cell in row.cells:
                if bg != "FFFFFF":
                    set_cell_background(cell, bg)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.line_spacing = 1.05
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    for r in p.runs:
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(9.0)
    if col_widths:
        for row in table.rows:
            for j, w in enumerate(col_widths):
                if j < len(row.cells):
                    row.cells[j].width = Inches(w)

def add_p(doc, text="", style='Normal', space_before=0, space_after=4, line_spacing=1.35, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = bold
        r.italic = italic
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.bold = True
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.bold = True
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.bold = True
    r.italic = True
    return p

def add_bullet(doc, lead_bold, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.left_indent = Inches(0.3)
    r_bold = p.add_run(lead_bold)
    r_bold.font.name = 'Times New Roman'
    r_bold.font.size = Pt(12)
    r_bold.bold = True
    r_text = p.add_run(text)
    r_text.font.name = 'Times New Roman'
    r_text.font.size = Pt(12)
    return p

def add_figure(doc, img_path, caption_text, width_in=5.8):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.keep_with_next = True
    p_img.add_run().add_picture(str(img_path), width=Inches(width_in))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.keep_with_next = False
    r = p_cap.add_run(caption_text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    r.bold = True
    return p_cap

def add_leader_line(doc, left_text, right_text, level=0, is_bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28 * level)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(2 if level == 0 else 1)
    p.paragraph_format.space_after = Pt(2 if level == 0 else 1)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.02), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    r1 = p.add_run(left_text)
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.bold = is_bold
    r_tab = p.add_run('\t')
    r2 = p.add_run(str(right_text))
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11)
    r2.bold = is_bold
    return p


def add_table_title(doc, title_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title_text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    r.bold = True
    return p

def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code_text)
    r.font.name = 'Courier New'
    r.font.size = Pt(9.5)
    return p
